#!/usr/bin/env python3
"""Offline fuzzy-match candidates: eval-run input products vs catalogue.

An offline approximation of the Power Query fuzzy merge (used while the PQ load inside the
workbook was still failing, and as a fallback / alternate-candidate source since). It is a
CANDIDATE list for manual one-to-one verification — not a verdict; Power Query's own
scores/order differ slightly (its engine adds character-level handling this proxy lacks).

Reads the cleaned workbook produced by build_catalogue_check.py:
  sheets 'Input products' (one title column) and 'Catalogue' (Sku Code | Product Name).

Writes (deterministic, no RNG):
  outputs/catalogue_check/fuzzy_match_offline_preview.csv
  outputs/catalogue_check/fuzzy_match_offline_preview.xlsx   (sorted, top-3 candidates)

Method
  * score = token-set Jaccard, tokens taken from a size-normalized string:
      - unicode fractions:  ½ → " 1/2", ¼ → " 1/4", ¾ → " 3/4"
      - glued mixed numbers: "11/2" → 1 1/2, "21/2" → 2 1/2
        (rule: 2-digit numerator, tens digit 1–3, units digit 1–4)
      - mixed forms "1-1/2" | "1.1/2" | "1 1/2" → ONE sentinel token  ~1_1_2
      - plain "1/2" → ~0_1_2
    WITHOUT the size normalization, 1 1/2-inch inputs mis-rank onto 1/2-inch decoys.
  * embedded-code pinning: a 7–9 digit run in an input title that matches a catalogue SKU
    is reported as-is (such codes are usually the mapping — still verify by eye).

Usage:
  python3 build_fuzzy_candidates.py [--workbook PATH] [--outdir DIR]
"""
import argparse
import csv
import re
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_WORKBOOK = ROOT / "outputs" / "catalogue_check" / "eval_inputs_vs_catalogue_vlookup.xlsx"
DEFAULT_OUTDIR = ROOT / "outputs" / "catalogue_check"

FRAC = {"\u00bd": " 1/2", "\u00bc": " 1/4", "\u00be": " 3/4"}
GLUED = re.compile(r"(?<![\d/])(\d{2})/(\d)(?![\d/])")
MIXED = re.compile(r"(\d+)\s*[-\. ]\s*(\d+)\s*/\s*(\d+)")
PLAIN_FRAC = re.compile(r"(?<![~\d])(\d+)/(\d+)")
TOKEN = re.compile(r"~[\d_]+|[a-z0-9]+")
EMBED = re.compile(r"(?<!\d)(\d{7,9})(?!\d)")


def norm_sizes(s: str) -> str:
    s = str(s)
    for k, v in FRAC.items():
        s = s.replace(k, v)

    def fix_glued(m: re.Match) -> str:
        b = m.group(1)
        if len(b) == 2 and b[0] in "123" and b[1] in "1234":
            return f"{b[0]} {b[1]}/{m.group(2)}"
        return m.group(0)

    s = GLUED.sub(fix_glued, s)
    s = MIXED.sub(lambda m: f" ~{m.group(1)}_{m.group(2)}_{m.group(3)} ", s)
    s = PLAIN_FRAC.sub(lambda m: f" ~0_{m.group(1)}_{m.group(2)} ", s)
    return s


def tok(s: str) -> frozenset:
    return frozenset(TOKEN.findall(norm_sizes(s).casefold()))


def jac(a: frozenset, b: frozenset) -> float:
    return len(a & b) / max(1, len(a | b))


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--workbook", default=str(DEFAULT_WORKBOOK), help="cleaned catalogue-check workbook")
    ap.add_argument("--outdir", default=str(DEFAULT_OUTDIR), help="output directory")
    args = ap.parse_args()
    outdir = Path(args.outdir)

    src = openpyxl.load_workbook(args.workbook, read_only=True, data_only=True)
    inputs = [r[0] for r in src["Input products"].iter_rows(min_row=2, values_only=True) if r[0]]
    cats = [(str(r[0]), r[1]) for r in src["Catalogue"].iter_rows(min_row=2, values_only=True) if r[1]]
    src.close()
    ctoks = [(tok(n), sku, n) for sku, n in cats]  # (tokens, sku, name)
    sku_set = {s for s, _ in cats}
    sku2name = {s: n for s, n in cats}

    rows = []
    for idx, inp in enumerate(inputs, 1):
        ti = tok(inp)
        scored = sorted(((jac(ti, tt), name, sku) for tt, sku, name in ctoks), key=lambda x: -x[0])
        emb = next(((c, sku2name[c]) for c in EMBED.findall(inp) if c in sku_set), ("", ""))
        rows.append((idx, inp, emb, scored[:3]))

    # ---- CSV (score-descending, stable) ----
    csv_path = outdir / "fuzzy_match_offline_preview.csv"
    with open(csv_path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["#", "input_product", "embedded_code", "code_found_in_catalogue",
                    "code_catalogue_name", "best1_name", "best1_sku", "best1_score",
                    "best2_name", "best2_sku", "best2_score", "best3_name", "best3_sku", "best3_score"])
        for idx, inp, emb, sc in sorted(rows, key=lambda x: -x[3][0][0]):
            r = [idx, inp, emb[0], "YES" if emb[0] else "", emb[1]]
            for t in sc:
                r += [t[1], t[2], round(t[0], 3)]
            w.writerow(r)

    # ---- XLSX ----
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Fuzzy candidates (offline)"
    hdr = ["#", "Input product (from eval run)", "Embedded code", "Code found in catalogue",
           "Code's catalogue name", "Suggested match", "SKU", "Approx score",
           "Alt 2", "Alt 2 SKU", "Alt 2 score", "Alt 3", "Alt 3 SKU", "Alt 3 score"]
    for j, h in enumerate(hdr, 1):
        c = ws.cell(row=1, column=j, value=h)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", start_color="DDEBF7")
    ws.freeze_panes = "A2"
    for idx, inp, emb, sc in sorted(rows, key=lambda x: -x[3][0][0]):
        r = [idx, inp, emb[0], "YES" if emb[0] else "", emb[1]]
        for t in sc:
            r += [t[1], t[2], round(t[0], 3)]
        ws.append(r)
    for col, wd in zip("BCEFH", [38, 58, 13, 58, 12]):
        ws.column_dimensions[col].width = wd
    xlsx_path = outdir / "fuzzy_match_offline_preview.xlsx"
    wb.save(xlsx_path)

    # ---- summary ----
    pin = sum(1 for r in rows if r[2][0])
    strong = sum(1 for r in rows if not r[2][0] and r[3][0][0] >= 0.75)
    med = sum(1 for r in rows if not r[2][0] and 0.60 <= r[3][0][0] < 0.75)
    weak = len(rows) - pin - strong - med
    print(f"rows: {len(rows)} | pinned(embedded code): {pin} | strong>=0.75: {strong} | "
          f"medium 0.60-0.75: {med} | weak<0.60: {weak}")
    print("csv: ", csv_path)
    print("xlsx:", xlsx_path)


if __name__ == "__main__":
    main()
