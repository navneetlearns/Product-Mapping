#!/usr/bin/env python3
"""Clean a sentinel dataset CSV: expand the line_item JSON, clean the product title.

Input : Tarun Sentinel Dataset.csv — one row per line item, where `line_item` is a JSON
        string {serial_no, hsn_sac_code, product_name, unit_rate_excl_gst}.
Output: "<stem> - cleaned.xlsx" + "<stem> - cleaned.csv" next to the source file.

What "clean" means here (same rules as the eval-mapping pipeline, `build_catalogue_check.clean`):
  * mojibake repair + zero-width/C1 strip + NBSP/whitespace collapse, then
  * the product title carries ONLY the title: an embedded SKU code is removed whether it is
    parenthesised ("... Birlanu (96000045)") or a bare leading token ("96000754 BEND 110MM..."),
    and any separator left dangling by that removal is trimmed.
  * brand words, dimensions, degree signs, quotes are untouched.

Every row that had a code in its title is flagged (`SKU in title?` = FLAG) and the extracted
code is checked against the row's mapped SKU.
"""
import argparse
import csv
import json
import re
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties

sys.path.insert(0, str(Path(__file__).resolve().parent))
from build_catalogue_check import clean  # mojibake/whitespace cleaner used by the eval builds

CODE_RE = re.compile(r"\d{8,9}")
PAREN_CODE_RE = re.compile(r"\s*\(\s*\d{8,9}\s*\)")
LEAD_CODE_RE = re.compile(r"^\s*\d{8,9}\s+")
TRAIL_SEP_RE = re.compile(r"[\s\-–,;/]+$")
DUP_SPACE_RE = re.compile(r"\s{2,}")


def clean_title(raw: str) -> str:
    t = clean(raw)
    t = PAREN_CODE_RE.sub("", t)
    t = LEAD_CODE_RE.sub("", t)
    t = TRAIL_SEP_RE.sub("", t)
    t = DUP_SPACE_RE.sub(" ", t).strip()
    t = re.sub(r"^[\s\-–,;/]+", "", t)
    return t.strip()


def load(src: Path):
    out = []
    with src.open(newline="", encoding="utf-8-sig") as fh:
        for r in csv.DictReader(fh):
            li = json.loads(r["line_item"]) if r["line_item"].strip() else {}
            eq = json.loads(r["equivalent_skus"]) if r["equivalent_skus"].strip() else []
            raw_title = li.get("product_name") or ""
            codes = CODE_RE.findall(raw_title)
            mapped = (r["sku"] or "").strip()
            out.append(
                {
                    "id": r["sentinel_item_id"],
                    "serial": li.get("serial_no"),
                    "hsn": li.get("hsn_sac_code"),
                    "rate": li.get("unit_rate_excl_gst"),
                    "title_raw": raw_title,
                    "title": clean_title(raw_title),
                    "codes": codes,
                    "has_sku": bool(codes),
                    "mapped": mapped,
                    "code_match": "yes" if codes and mapped in codes else ("no" if codes else ""),
                    "eq": eq,
                    "exclusive": (r["distributor_exclusive"] or "").strip().lower() == "true",
                    "invoice": (r["source_invoice"] or "").strip(),
                    "line_item_raw": r["line_item"],
                }
            )
    return out


HEAD_FILL = PatternFill("solid", fgColor="14342A")
HEAD_FONT = Font(color="FFFFFF", bold=True, size=10, name="Aptos")
FLAG_FILL = PatternFill("solid", fgColor="FCE9C8")
MONO = Font(name="Consolas", size=10)


def style_header(ws, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(ncols)}{ws.max_row}"


def build(src: Path):
    rows = load(src)
    flagged = [r for r in rows if r["has_sku"]]

    by_sku = {}
    for r in rows:
        by_sku.setdefault(r["mapped"], []).append(r)
    dup_sku = {k: v for k, v in by_sku.items() if len(v) > 1}

    wb = openpyxl.Workbook()

    # ---- sheet 1: input products (line_item expanded) ----
    ws = wb.active
    ws.title = "Input products"
    cols = ["#", "sentinel_item_id", "serial_no", "product_title", "SKU in title?",
            "code found in title", "mapped sku", "code = mapped sku", "hsn_sac_code",
            "unit_rate_excl_gst", "equivalent skus", "source_invoice"]
    ws.append(cols)
    for i, r in enumerate(rows, 1):
        ws.append([i, r["id"], r["serial"], r["title"], "FLAG" if r["has_sku"] else "",
                   ", ".join(r["codes"]), r["mapped"], r["code_match"],
                   r["hsn"], r["rate"], ", ".join(r["eq"]), r["invoice"]])
        if r["has_sku"]:
            for c in range(1, len(cols) + 1):
                ws.cell(row=ws.max_row, column=c).fill = FLAG_FILL
        ws.cell(row=ws.max_row, column=9).number_format = "@"  # HSN stays text
        ws.cell(row=ws.max_row, column=10).number_format = "#,##0"
    style_header(ws, len(cols))
    widths = [5, 16, 9, 62, 14, 20, 12, 16, 13, 16, 32, 42]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2, min_col=7, max_col=9):
        for cell in row:
            cell.font = MONO

    # ---- sheet 2: flagged (raw title kept for audit) — only when something is flagged ----
    if flagged:
        ws2 = wb.create_sheet("Flagged titles")
        cols2 = ["sentinel_item_id", "product_title (cleaned)", "original title", "how",
                 "code found", "mapped sku", "code = mapped sku"]
        ws2.append(cols2)
        for r in flagged:
            how = "parenthesised" if PAREN_CODE_RE.search(r["title_raw"]) else "leading"
            ws2.append([r["id"], r["title"], r["title_raw"], how,
                        ", ".join(r["codes"]), r["mapped"], r["code_match"]])
        style_header(ws2, len(cols2))
        for i, w in enumerate([16, 58, 58, 14, 18, 12, 14], 1):
            ws2.column_dimensions[get_column_letter(i)].width = w
        for row in ws2.iter_rows(min_row=2, min_col=4, max_col=4):
            row[0].alignment = Alignment(horizontal="center")

    # ---- sheet 3: same sku, different titles (audit) ----
    ws3 = wb.create_sheet("Same SKU, many titles")
    ws3.append(["mapped sku", "titles", "sentinel_item_ids", "titles (cleaned)"])
    for sku, v in sorted(dup_sku.items(), key=lambda kv: (-len(kv[1]), kv[0])):
        ws3.append([sku, len(v), ", ".join(x["id"] for x in v),
                    " || ".join(x["title"] for x in v)])
    style_header(ws3, 4)
    for i, w in enumerate([12, 8, 34, 110], 1):
        ws3.column_dimensions[get_column_letter(i)].width = w

    # ---- sheet 4: summary ----
    ws4 = wb.create_sheet("Summary")
    n = len(rows)
    summary = [
        ("source file", src.name),
        ("line items (rows)", n),
        ("titles carrying a SKU code", len(flagged)),
        ("  - parenthesised  e.g. 'Birlanu (96000045)'",
         sum(1 for r in flagged if PAREN_CODE_RE.search(r["title_raw"]))),
        ("  - leading bare   e.g. '96000754 BEND 110MM...'",
         sum(1 for r in flagged if LEAD_CODE_RE.match(r["title_raw"]))),
        ("code = mapped sku", sum(1 for r in flagged if r["code_match"] == "yes")),
        ("code != mapped sku", sum(1 for r in flagged if r["code_match"] == "no")),
        ("titles changed beyond code removal (whitespace/mojibake)",
         sum(1 for r in rows if clean(r["title_raw"]) != r["title_raw"])),
        ("distinct mapped skus", len(by_sku)),
        ("skus appearing on more than one title", len(dup_sku)),
        ("rows with equivalent_skus", sum(1 for r in rows if r["eq"])),
        ("distributor_exclusive = true", sum(1 for r in rows if r["exclusive"])),
        ("hsn 8-digit", sum(1 for r in rows if r["hsn"] and len(str(r["hsn"])) == 8)),
        ("hsn 4-digit (family only)", sum(1 for r in rows if r["hsn"] and len(str(r["hsn"])) == 4)),
        ("hsn missing", sum(1 for r in rows if not r["hsn"])),
    ]
    ws4.append(["metric", "value"])
    for k, v in summary:
        ws4.append([k, v])
    style_header(ws4, 2)
    ws4.column_dimensions["A"].width = 54
    ws4.column_dimensions["B"].width = 60
    for row in ws4.iter_rows(min_row=2, min_col=2, max_col=2):
        row[0].alignment = Alignment(horizontal="left")

    # ---- sheet 5: raw line_item (audit) ----
    ws5 = wb.create_sheet("Raw line_item")
    ws5.append(["sentinel_item_id", "line_item (original JSON)"])
    for r in rows:
        ws5.append([r["id"], r["line_item_raw"]])
    style_header(ws5, 2)
    ws5.column_dimensions["A"].width = 18
    ws5.column_dimensions["B"].width = 120

    # print setup: landscape, one page wide, header row repeated
    for sheet in wb.worksheets:
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        sheet.print_title_rows = "1:1"

    out_xlsx = src.with_name(src.stem + " - cleaned.xlsx")
    wb.save(out_xlsx)

    out_csv = src.with_name(src.stem + " - cleaned.csv")
    with out_csv.open("w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["sentinel_item_id", "serial_no", "product_title", "sku_in_title",
                    "code_found_in_title", "mapped_sku", "code_matches_mapped_sku",
                    "hsn_sac_code", "unit_rate_excl_gst", "equivalent_skus",
                    "distributor_exclusive", "source_invoice"])
        for r in rows:
            w.writerow([r["id"], r["serial"], r["title"], "FLAG" if r["has_sku"] else "",
                        ",".join(r["codes"]), r["mapped"], r["code_match"], r["hsn"],
                        r["rate"], ",".join(r["eq"]), str(r["exclusive"]).lower(), r["invoice"]])
    return out_xlsx, out_csv, rows, flagged, dup_sku


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("src")
    a = ap.parse_args()
    x, c, rows, flagged, dup = build(Path(a.src))
    print(f"wrote {x}\nwrote {c}")
    print(f"rows={len(rows)} flagged={len(flagged)} dup_sku_groups={len(dup)}")
