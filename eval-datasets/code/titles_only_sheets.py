#!/usr/bin/env python3
"""Write one workbook with ONE column per lane: cleaned product titles only.

Usage:
  python3 titles_only_sheets.py --out <xlsx> "<lane label>=<csv>" ["<lane label>=<csv>" ...]

Each lane becomes a sheet whose only content is a `product_title` header plus the cleaned
titles — no ids, skus, HSN, rates, invoices, flags or extra sheets.
"""
import argparse
import csv
import json
import re

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill


def clean_title(raw):
    t = " ".join(str(raw or "").split())  # keep the source apostrophes: ’ vs ’’ encodes ft vs inch
    t = re.sub(r"\s*\(\s*\d{8,9}\s*\)", "", t)
    t = re.sub(r"^\s*\d{8,9}\s+", "", t)
    return re.sub(r"[\s\-–,;/]+$", "", t).strip()


def titles(csv_path):
    out = []
    with open(csv_path, newline="", encoding="utf-8-sig") as fh:
        for r in csv.DictReader(fh):
            li = json.loads(r["line_item"]) if r["line_item"].strip() else {}
            out.append(clean_title(li.get("product_name")))
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", required=True)
    ap.add_argument("lanes", nargs="+", help='label=csv')
    a = ap.parse_args()

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for spec in a.lanes:
        label, path = spec.split("=", 1)
        ws = wb.create_sheet(label[:31])
        ws.append(["product_title"])
        for t in titles(path):
            ws.append([t])
        c = ws.cell(row=1, column=1)
        c.font = Font(bold=True, color="FFFFFF", size=10, name="Aptos")
        c.fill = PatternFill("solid", fgColor="14342A")
        c.alignment = Alignment(vertical="center")
        ws.column_dimensions["A"].width = 64
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:A{ws.max_row}"
        ws.page_setup.orientation = "portrait"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        print(f"{label}: {ws.max_row - 1} titles")
    wb.save(a.out)
    print("wrote", a.out)


if __name__ == "__main__":
    main()
