#!/usr/bin/env python3
"""Resolve the 42 open catalogue-check cases from title text (no SKU codes in titles).

Context: of the 100 brand-name eval-run input products, 58 are located (48 Power Query
fuzzy matches + 10 code-rescued blanks); 42 remain — no embedded code, fuzzy score < 0.70.
This builder produces a MANUAL-VERIFICATION workbook for those 42:

  * a title parser (normalization v2 + synonym/class tables + mm<->inch size canon),
  * candidate generation per input (type + size hard gates; class/length/material scored),
  * the run export's recorded GT SKU as a "lead" column (all leads verified in-export),
  * per-component flags + variant-gap notes, and an EMPTY verdict column for the user.

Strictness: per user decision (2026-10-19), "found" requires an EXACT variant — the
workbook flags AMBIG-VARIANT where catalogue variants differ on a dimension the title
does not pin (length / class / colour / fit type). The script NEVER decides a match.

Sources (same folder-only scope as the rest of the check):
  - workbook  : F:/.../Eval Related Product Mapping/eval_inputs_vs_catalogue_vlookup.xlsx
                (sheet 'Input Products (2)' — blank score => open case; no 8-digit code => in scope)
  - run       : F:/.../turnover-product-mapping-c4cce19d.csv (GT via reference_outputs)
  - catalogue : F:/.../products-18-October-2026-10-52-16.csv (cleaned with the same rules)

Outputs (repo):
  outputs/catalogue_check/open_42_resolution.xlsx
  outputs/catalogue_check/open_42_resolution.csv
  outputs/catalogue_check/open_42_candidates_long.csv
  outputs/catalogue_check/open_cases_42.csv        (manifest of the 42 inputs + GT leads)
With --copy-to <dir>, the xlsx is copied there (default: the F: drop folder).

Usage:
  python3 build_open_case_resolution.py [--workbook PATH] [--run PATH] [--catalogue PATH]
                                        [--outdir DIR] [--copy-to DIR]
"""
from __future__ import annotations

import argparse
import csv
import re
import shutil
import sys
import tempfile
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from build_catalogue_check import clean  # same cleaning path as the rest of the check

HERE = Path(__file__).resolve().parent
OUTDIR = HERE.parent / "outputs" / "catalogue_check"

DROP = Path("/mnt/f/Product/zotok/Product Mapping/Eval Related Product Mapping")
DEF_WB = DROP / "eval_inputs_vs_catalogue_vlookup.xlsx"
DEF_RUN = DROP / "turnover-product-mapping-c4cce19d.csv"
DEF_CAT = DROP / "products-18-October-2026-10-52-16.csv"

CODE_RE = re.compile(r"(?<!\d)\d{8}(?!\d)")

# --------------------------------------------------------------------------- #
# Size canon: mm <-> inch trade equivalents.  Two conventions coexist in this
# catalogue: CPVC/UPVC CTS-style fittings (2" = 50 mm) and PVC pressure/SWR
# metric pipes (2" = 63 mm, 4" = 110 mm).  An inch token therefore maps to a
# SET of possible mm values; a metric token maps to itself.
# --------------------------------------------------------------------------- #
INCH_MM = {
    "1/2": {15}, "3/4": {20}, "1": {25}, "1 1/4": {32}, "1 1/2": {40},
    "2": {50, 63}, "2 1/2": {75}, "3": {90}, "4": {110}, "5": {140}, "6": {160},
}
FRAC_WORDS = {"1/2": 15, "3/4": 20}  # kept for readability in debug output

TYPE_PATTERNS = {
    "RED_TEE":       r"RED(?:UCING|UCG|UC)?[.\s]*TEE",
    "RED_BUSH":      r"REDUCING\s*BUSH|RED\.?\s*BUSH",
    "RED_SOCKET":    r"REDUCING\s*(?:SOCKET|COUPLER)",
    "RED_ELBOW":     r"REDUCING\s*ELBOW",
    "REDUCING_GEN":  r"REDUCING\s*(?:SOCKET|ELBOW|BUSH|TEE)|RED[.\s]*TEE",
    "SOCKET":        r"\bSOCKET\b|\bCOUPLER\b|\bSOC\b",
    "END_CAP":       r"END\s*CAP",
    "WATER_TANK":    r"WATER\s*TANK",
    "SUBMERSIBLE":   r"SUBMERSIBLE",
    "MTA":           r"\bMTA\b",
    "FTA":           r"\bFTA\b",
    "TEE":           r"\bTEE\b",
    "PIPE":          r"\bPIPES?\b",
    "ELBOW":         r"ELBOW|\bELBO\b|BEND",
    "NIPPLE":        r"\bNIPPLE\b",
}

SUBTYPE_LEX = ["FAB", "TRD", "DOOR", "SOLFIT", "PUSHFIT", "PLASTICS", "TRANS",
               "LW", "PLAIN", "SILENT", "UGD", "MEDIUM", "STANDARD", "HEAVY",
               "SAMPLE", "SAMPL", "WHITE", "GREEN", "YELLOW", "BLUE", "BLACK",
               "V4-12.5", "S/F", "S/S", "D/S", "BRASS", "THREADED"]

# kg/cm^2 pressure class <-> catalogue CLASS-N naming (CLASS-1=2.5kg and
# 4kg=>LW 4Kg are GT-verified; the rest validated by the GT unit test below).
KG_CLASS = {2.5: "1", 4: "2", 6: "3", 10: None, 25: None}
CLASS_KG = {"1": 2.5, "2": 4, "3": 6}

UNIT_SPACE = re.compile(r"(?<=\d)\s*(MM|KG|MTRS?|MRTS?|METER|METERS|FT|SCH|SDR|LITRES?|LITERS?|LTRS?)\b")
X_SPACE = re.compile(r"(?<=[A-Z0-9])\s*X\s*(?=\d)")
GLUED = re.compile(r"(?<![\d/])(\d{2})/(\d)(?![\d/])")
PAREN_LEN = re.compile(r"\(\s*(\d{3})\s*MM\s*\)|\(\s*(\d)\s*M(?:TR|ETER|RT)?S?\s*\)")
INCH_TOK = re.compile(r'(\d+\s+\d/\d|\d+/\d|\d+)\s*"')
INCH_WORD = re.compile(r"(\d+)\s*INCH(?:I|ES)?\b")
BARE_FRAC = re.compile(r"(?<![\d./])(\d+/\d+)(?![\d/])")
DEC_INCH = re.compile(r"(?<![\d.])(0\.5|0\.75|1\.25|1\.5|2\.0)(?!\d)(?!\s*(?:KG|MM|M\b|FT|L\b))")
DUAL_INCH_RAW = re.compile(r"(?<![\d./])(\d+/\d+|\d{1,2}(?:\.\d{1,2})?)\s+X\s+(\d+/\d+|\d{1,2}(?:\.\d{1,2})?)(?!\s*(?:KG|MM|M\b|FT|L\b))")
DEC_MAP = {"0.5": "1/2", "0.75": "3/4", "1.25": "1 1/4", "1.5": "1 1/2", "2.0": "2"}
MM_TOK = re.compile(r"(\d{2,3})\s*MM(?![A-Z])")
DUAL_MM = re.compile(r"(\d{2,3})\s*X\s*(\d{2,3})\s*MM")
LEN_TOK = re.compile(r"(\d{1,2})\s*(MTRS?|MRTS?|METER|METERS|M)\.?(?![A-Z])|(\d{1,2})\s*FT\b")
CLASS_TOK = re.compile(r"SCH\s*[-.]?\s*(\d{2})|(\d{2})\s*SCH\b|SDR\s*[-.]?\s*(\d+(?:\.\d+)?)")
SCH_S = re.compile(r"SCH\s*\(\s*S\s*(\d{2})\s*\)|\bS(40|80)\b")
KG_TOK = re.compile(r"(\d+(?:\.\d+)?)\s*KG(?:\s*/\s*CM2)?")
CLASS_N = re.compile(r"CLASS\s*-?\s*(\d)")
PN_TOK = re.compile(r"\bPN\s*(\d+)\b")
TYPE_LET = re.compile(r"TYPE\s*-?\s*([AB])\b")
LAYER_TOK = re.compile(r"(\d)\s*LAYER|\b([34])\s*L(?!ITRE|ITER|TR)\b")
CAP_TOK = re.compile(r"(\d{3,4})\s*(?:LITRES?|LITERS?|L)\b")
DEG45 = re.compile(r"45")
ELBOW_WORD = re.compile(r"ELBOW|\bELBO\b|BEND|SHOE")


def norm(s: str) -> str:
    s = clean(str(s)).upper()
    s = s.replace("\u2019\u2019", ' "')                  # ’’ = inches (before single ’ = ft)
    s = s.replace("\u2033", '"').replace("\u201d", '"').replace("\u201c", '"')
    s = re.sub(r"(?<=\d)\s*\u2019", " FT", s)          # 20’ -> 20 FT
    s = re.sub(r"(?<=\d)\s*''", ' "', s)               # 1 1/2'' -> inches
    s = re.sub(r"(?<=\d)\s*'(?!')", ' "', s)           # lone apostrophe as inch
    s = X_SPACE.sub(" X ", s)
    s = UNIT_SPACE.sub(r" \1", s)

    def glue(m):
        b = m.group(1)
        if len(b) == 2 and b[0] in "123" and b[1] in "1234":
            return f"{b[0]} {b[1]}/{m.group(2)}"
        return m.group(0)

    return GLUED.sub(glue, s)


def parse(s: str) -> dict:
    """Structured components of a title (input or catalogue name)."""
    t = norm(s)
    lengths, classes, sizes = set(), set(), []
    subtype = {tok for tok in SUBTYPE_LEX if re.search(re.escape(tok), t)}
    material = None
    for m, name in (("CPVC", "CPVC"), ("UPVC", "UPVC"), ("SWR", "SWR"), ("AGRI", "AGRI")):
        if m in t:
            material = name
            break
    if material is None and re.search(r"\bPVC\b", t):
        material = "PVC"

    body = t

    def take(rx, fn):
        nonlocal body
        for m in rx.finditer(body):
            fn(m)
        body = rx.sub(" ", body)

    def paren(m):
        if m.group(1):
            lengths.add("L300MM")
        else:
            lengths.add(f"L{m.group(2)}M")

    def dual_mm(m):
        sizes.append({int(m.group(1))})
        sizes.append({int(m.group(2))})

    def mm(m):
        sizes.append({int(m.group(1))})

    def dual_inch_raw(m):
        for raw in (m.group(1), m.group(2)):
            key = DEC_MAP.get(raw, raw)
            if key in INCH_MM:
                sizes.append(set(INCH_MM[key]))

    def inch(m):
        raw = " ".join(m.group(1).split())
        if raw in INCH_MM:
            sizes.append(set(INCH_MM[raw]))

    def inch_word(m):
        if m.group(1) in INCH_MM:
            sizes.append(set(INCH_MM[m.group(1)]))

    def mixed_bare(m):
        raw = f"{m.group(1)} {m.group(2)}/{m.group(3)}"
        if raw in INCH_MM:
            sizes.append(set(INCH_MM[raw]))

    def bare_frac(m):
        if m.group(1) in INCH_MM:
            sizes.append(set(INCH_MM[m.group(1)]))

    def dec_inch(m):
        key = DEC_MAP.get(m.group(1))
        if key and key in INCH_MM:
            sizes.append(set(INCH_MM[key]))

    def length(m):
        if m.group(1):
            lengths.add(f"L{m.group(1)}M")
        else:
            lengths.add(f"L{m.group(3)}FT")

    def sch(m):
        if m.group(1):
            classes.add(f"SCH{m.group(1)}")
        elif m.group(2):
            classes.add(f"SCH{m.group(2)}")
        elif m.group(3):
            classes.add(f"SDR{m.group(3)}")

    def sch_s(m):
        classes.add(f"SCH{m.group(1) or m.group(2)}")

    def kg(m):
        classes.add(f"KG{float(m.group(1)):g}")

    def cls_n(m):
        classes.add(f"CLASS{m.group(1)}")

    def pn(m):
        classes.add(f"PN{m.group(1)}")

    def type_let(m):
        classes.add(f"TYPE{m.group(1)}")

    def layer(m):
        classes.add(f"LAYER{m.group(1) or m.group(2)}")

    def cap(m):
        classes.add(f"CAP{m.group(1)}")

    take(PAREN_LEN, paren)
    take(DUAL_MM, dual_mm)
    take(MM_TOK, mm)
    take(DUAL_INCH_RAW, dual_inch_raw)
    take(INCH_TOK, inch)
    take(INCH_WORD, inch_word)
    take(re.compile(r"(?<![\d.])(\d+)\s+(\d)/(\d)(?![\d/])"), mixed_bare)
    take(BARE_FRAC, bare_frac)
    take(DEC_INCH, dec_inch)
    take(LEN_TOK, length)
    take(CLASS_TOK, sch)
    take(SCH_S, sch_s)
    take(KG_TOK, kg)
    take(CLASS_N, cls_n)
    take(PN_TOK, pn)
    take(TYPE_LET, type_let)
    take(LAYER_TOK, layer)
    take(CAP_TOK, cap)

    return {"norm": t, "sizes": sizes, "lengths": lengths, "classes": classes,
            "subtype": subtype, "material": material,
            "angle": ("45" if (ELBOW_WORD.search(t) and re.search(r"\b45\b|45\s*[CD]|45DEG", t))
                      else ("90" if (ELBOW_WORD.search(t) and re.search(r"\b90\b|90\s*[CD]|90DEG", t)) else ""))}


def canon_classes(classes: set) -> set:
    """Collapse class naming equivalents (CLASS-2 <-> KG4) for variant-gap display."""
    out = set()
    for c in classes:
        if c.startswith("CLASS") and c[5:] in CLASS_KG:
            out.add(f"KG{CLASS_KG[c[5:]]:g}")
        else:
            out.add(c)
    return out


def input_concept(t: str) -> str | None:
    """Primary product-type concept of an INPUT title (ordered rules)."""
    if re.search(r"RED\.?\s*TEE|REDUCING\s+TEE", t):
        return "RED_TEE"
    if re.search(r"REDUCING\s+BUSH", t):
        return "RED_BUSH"
    if re.search(r"REDUCING\s+(?:COUPLER|SOCKET)", t):
        return "RED_SOCKET"
    if re.search(r"REDUCING\s+ELBOW", t):
        return "RED_ELBOW"
    if re.search(r"\bREDUCER\b|REDUCING\b", t):
        return "REDUCING_GEN"
    if re.search(r"COUPLER", t):
        return "SOCKET"
    if re.search(r"END\s*CAP", t):
        return "END_CAP"
    if re.search(r"TANK", t) and not re.search(r"TANK\s*NIPPLE", t):
        return "WATER_TANK"
    if re.search(r"COLUMN", t):
        return "SUBMERSIBLE"
    if re.search(r"\bMTA\b", t):
        return "MTA"
    if re.search(r"\bFTA\b", t):
        return "FTA"
    if re.search(r"SHOE", t):
        return "ELBOW45"
    if ELBOW_WORD.search(t):
        return "ELBOW45" if DEG45.search(t) else "ELBOW"
    if re.search(r"\bTEE\b", t):
        return "TEE"
    if re.search(r"HALFSET", t):
        return "UNKNOWN"
    if re.search(r"\bSCH\b", t):
        return "PIPE"
    if re.search(r"\bPIPES?\b", t):
        return "PIPE"
    return None


def candidate_type_match(concept: str | None, ct: str, cp: dict) -> bool:
    if concept is None:
        return True
    if concept == "ELBOW45":
        return bool(ELBOW_WORD.search(ct)) and bool(DEG45.search(ct))
    if concept == "UNKNOWN":
        return True
    if concept == "PIPE" and re.search(r"CLAMP|\bCLIP\b|BRACKET|PLUNGER|HANGER|SOLVENT|ADHESIVE", ct):
        return False
    pat = TYPE_PATTERNS.get(concept)
    return bool(re.search(pat, ct)) if pat else True


def size_ok(inp_sizes, cand_sizes) -> bool:
    if not inp_sizes:
        return True
    union = set()
    for cs in cand_sizes:
        union |= cs
    return all(bool(iset & union) for iset in inp_sizes)


def class_ok(inp_classes, cand_classes, ct) -> bool:
    if not inp_classes:
        return True
    for c in inp_classes:
        hit = False
        if c.startswith("KG"):
            kg = float(c[2:])
            cls = KG_CLASS.get(kg)
            if re.search(rf"{re.escape(format(kg, 'g'))}\s*KG", ct):
                hit = True
            if cls and f"CLASS{cls}" in cand_classes:
                hit = True
            if cls and re.search(rf"CLASS\s*-?\s*{cls}\b", ct):
                hit = True
            if kg == 25 and re.search(r"25\s*KG", ct):
                hit = True
        elif c.startswith("CLASS"):
            n = c[5:]
            if f"CLASS{n}" in cand_classes or re.search(rf"CLASS\s*-?\s*{n}\b", ct):
                hit = True
            kg = CLASS_KG.get(n)
            if kg and re.search(rf"{kg:g}\s*KG", ct):
                hit = True
        elif c.startswith("SCH"):
            n = c[3:]
            if re.search(rf"SCH\s*[-.]?\s*{n}\b", ct):
                hit = True
        elif c.startswith("SDR"):
            n = c[3:]
            if re.search(rf"SDR\s*[-.]?\s*{re.escape(n)}\b", ct):
                hit = True
        elif c.startswith("PN"):
            if re.search(rf"PN\s*{c[2:]}\b", ct):
                hit = True
        elif c.startswith("TYPE"):
            if re.search(rf"TYPE\s*-?\s*{c[4:]}\b", ct):
                hit = True
        elif c.startswith("LAYER"):
            if re.search(rf"{c[5:]}\s*LAYER", ct):
                hit = True
        elif c.startswith("CAP"):
            if re.search(rf"{c[3:]}\s*(?:LITRES?|LITERS?|L)\b", ct):
                hit = True
        if not hit:
            return False
    return True


def length_ok(inp_lengths, cand_lengths) -> bool:
    return (not inp_lengths) or bool(inp_lengths & cand_lengths)


def material_score(inp_material, cand_material) -> int:
    if not inp_material:
        return 0
    if inp_material == cand_material:
        return 2
    if inp_material == "PVC" and cand_material in (None, "PVC"):
        return 1
    if inp_material == "AGRI" and cand_material in ("PVC", None):
        return 1
    return 0


def jaccard(a: set, b: set) -> float:
    return len(a & b) / max(1, len(a | b))


def esc(s: str) -> str:
    return re.escape(s)


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--workbook", default=str(DEF_WB))
    ap.add_argument("--run", default=str(DEF_RUN))
    ap.add_argument("--catalogue", default=str(DEF_CAT))
    ap.add_argument("--outdir", default=str(OUTDIR))
    ap.add_argument("--copy-to", default=str(DROP))
    args = ap.parse_args()
    outdir = Path(args.outdir)
    outdir.mkdir(parents=True, exist_ok=True)

    # ---------------- the 42 (from the saved PQ workbook) ----------------
    tmp = Path(tempfile.mkdtemp()) / "_wb.xlsx"
    shutil.copy2(args.workbook, tmp)          # works even while Excel holds the file
    wb = openpyxl.load_workbook(tmp, data_only=True)
    ws = wb["Input Products (2)"]
    hdr = [c.value for c in ws[1]]
    assert hdr[0] and "Input" in str(hdr[0]), hdr
    blanks = [r[0] for r in ws.iter_rows(min_row=2, values_only=True) if r[0] and not r[3]]
    wb.close()
    blanks = [clean(t) for t in blanks]
    with_code = [t for t in blanks if CODE_RE.search(t)]
    the42 = [t for t in blanks if not CODE_RE.search(t)]
    print(f"blank rows: {len(blanks)} | code-carrying (rescued earlier): {len(with_code)} | open cases: {len(the42)}")
    assert len(the42) == 42, f"expected 42 open cases, found {len(the42)}"

    # ---------------- catalogue ----------------
    raw = list(csv.DictReader(open(args.catalogue, newline="", encoding="utf-8")))
    cat, seen = [], set()
    for r in raw:
        sku, name = clean(r.get("Sku Code")), clean(r.get("Product Name"))
        if sku in seen:
            continue
        seen.add(sku)
        cat.append((sku, name))
    cat.sort(key=lambda t: (len(t[0]), t[0]))
    sku2name = dict(cat)
    print(f"catalogue: {len(cat)} products")

    # ---------------- run export: GT leads ----------------
    rows = list(csv.DictReader(open(args.run, newline="", encoding="utf-8")))
    gt = {}
    for r in rows:
        import json as _json
        inp = _json.loads(r["inputs"])
        title = clean(inp["line_items"][0]["product_name"])
        if title in gt:
            continue
        ref = _json.loads(r["reference_outputs"]) if r.get("reference_outputs") else {}
        exp = (ref.get("expected") or [{}])[0]
        gt[title] = {"status": exp.get("expected_status", ""), "sku": exp.get("expected_sku", "") or ""}
    print(f"run rows: {len(rows)} -> {len(gt)} unique titles")
    missing_gt = [t for t in the42 if t not in gt]
    assert not missing_gt, f"titles missing from run export: {missing_gt}"

    leads = {t: gt[t] for t in the42 if gt[t]["sku"]}
    leads_in = {t: (g["sku"] in sku2name) for t, g in leads.items()}
    print(f"GT leads among the 42: {len(leads)} | all in export: {all(leads_in.values())}")

    # ---------------- candidate generation ----------------
    catp = [(sku, name, parse(name)) for sku, name in cat]
    results, cand_long, no_match = [], [], []
    for idx, title in enumerate(the42, 1):
        ip = parse(title)
        concept = input_concept(ip["norm"])
        itokens = set(re.findall(r"[A-Z0-9/]+", ip["norm"]))
        matched = []
        for sku, name, cp in catp:
            ct = norm(name)
            if not candidate_type_match(concept, ct, cp):
                continue
            if not size_ok(ip["sizes"], cp["sizes"]):
                continue
            flags = {
                "type": True,
                "size": bool(ip["sizes"]),
                "class": class_ok(ip["classes"], cp["classes"], ct),
                "length": length_ok(ip["lengths"], cp["lengths"]),
                "material": material_score(ip["material"], cp["material"]),
            }
            score = (2 * flags["type"] + 2 * (1 if ip["sizes"] else 0)
                     + 2 * flags["class"] + 1.5 * flags["length"]
                     + flags["material"]
                     + 0.75 * len(ip["subtype"] & cp["subtype"])
                     + jaccard(itokens, set(re.findall(r"[A-Z0-9/]+", ct))))
            matched.append({"sku": sku, "name": name, "flags": flags, "score": round(score, 3),
                            "cls": cp["classes"], "len": cp["lengths"], "sub": cp["subtype"],
                            "mat": cp["material"], "angle": cp["angle"]})
        matched.sort(key=lambda m: (-m["score"], m["sku"]))
        # strict pool: every dimension the title specifies must match
        strict = [m for m in matched
                  if m["flags"]["class"] and m["flags"]["length"]
                  and ((not ip["material"]) or m["flags"]["material"] > 0)]
        top = (strict + [m for m in matched if m not in strict])[:5]
        best = (strict or matched)[0] if matched else None

        # variant gaps over the strictest pool; only dims the title leaves open
        gaps, pool = [], (strict or matched)[:6]
        dims = []
        if not ip["lengths"]:
            dims.append(("length", lambda m: tuple(sorted(m["len"]))))
        if not ip["classes"]:
            dims.append(("class", lambda m: tuple(sorted(canon_classes(m["cls"])))))
        if not ip["angle"]:
            dims.append(("angle", lambda m: (m["angle"] or "–",)))
        if not ip["material"]:
            dims.append(("material", lambda m: (m["mat"] or "–",)))
        dims.append(("subtype", lambda m: tuple(sorted(m["sub"]))))
        if pool:
            for dim, get in dims:
                vals = {get(m) for m in pool}
                if len(vals) > 1:
                    pretty = " | ".join("/".join(str(x) for x in v) or "–"
                                        for v in sorted(vals, key=str))
                    gaps.append(f"{dim}: {pretty}")

        def fl(v, key):
            if key == "type":
                return "Y"
            if key == "size":
                return "Y" if ip["sizes"] else "–"
            if key == "material":
                if not ip["material"]:
                    return "–"
                return "Y" if v else "N"
            return "Y" if v else "N"

        gt_sku = gt[title]["sku"]
        gt_status = gt[title]["status"]
        if gt_sku:
            suggestion = "GT-LEAD"
        elif not matched:
            suggestion = "NO-MATCH"
        elif len(strict) == 1:
            suggestion = "EXACT candidate"
        else:
            suggestion = "AMBIG-VARIANT"

        notes = []
        if concept == "UNKNOWN":
            notes.append("type word not in catalogue vocabulary")
        if gt_sku:
            gt_rank = next((i for i, m in enumerate(matched, 1) if m["sku"] == gt_sku), None)
            if gt_rank is None:
                notes.append("GT lead not among type+size candidates")
            else:
                notes.append(f"GT lead is candidate #{gt_rank} of {len(matched)} by component score")
        if ip["lengths"] and not any(ip["lengths"] & m["len"] for m in matched):
            notes.append(f"title length {'/'.join(sorted(ip['lengths']))} not present in candidates")
        if ip["classes"] and not any(m["flags"]["class"] for m in matched):
            notes.append(f"title class {'/'.join(sorted(ip['classes']))} not satisfied by any candidate")
        if ip["material"] and not any(m["flags"]["material"] > 0 for m in matched):
            notes.append(f"material {ip['material']} not in candidate names")
        if not strict and matched:
            notes.append("no candidate matched all specified dimensions")
        if len(strict) > 1 and not gaps:
            notes.append(f"{len(strict)} candidates identical on all checked dimensions")
        if concept == "ELBOW45":
            notes.append("45-deg elbow family (shoe=45 synonym)")
        if ip["material"] == "AGRI":
            notes.append("agri -> Pressure Pipe CLASS-N / AGRI PN")
        if "BRASS" in ip["norm"]:
            notes.append("BRASS token absent from export names")
        note_str = "; ".join(notes)

        results.append({
            "idx": idx, "title": title, "concept": concept, "gt_sku": gt_sku,
            "gt_name": sku2name.get(gt_sku, "") if gt_sku else "",
            "gt_in": ("Y" if gt_sku in sku2name else ("-" if not gt_sku else "NO")),
            "gt_status": gt_status, "suggestion": suggestion,
            "best": best, "top": top, "gaps": gaps, "notes": notes, "note_str": note_str,
            "flags": ({k: fl(v, k) for k, v in best["flags"].items()} if best else {}),
            "material": ip["material"], "classes": sorted(ip["classes"]),
            "lengths": sorted(ip["lengths"]), "sizes": ip["sizes"],
        })
        for rank, m in enumerate(top, 1):
            cand_long.append({"idx": idx, "title": title, "rank": rank, "sku": m["sku"],
                              "name": m["name"], "score": m["score"],
                              "type": "Y", "size": "Y" if ip["sizes"] else "–",
                              "class": "Y" if m["flags"]["class"] else "N",
                              "length": "Y" if m["flags"]["length"] else "N",
                              "material": m["mat"] or "–"})
        if suggestion == "NO-MATCH" or concept in ("UNKNOWN", None):
            cand_tokens = [tok for tok in re.findall(r"[A-Z]{3,}", ip["norm"])]
            zero = []
            for tok in dict.fromkeys(cand_tokens):
                if tok in ("BIRLA", "BIRLANU", "HIL", "AND", "THE"):
                    continue
                if not any(tok in norm(n) for _, n in cat):
                    zero.append(tok)
            no_match.append({"idx": idx, "title": title, "zero_tokens": zero,
                             "concept": concept})

    # ---------------- QA: GT unit test ----------------
    top3_hits, misses = 0, []
    for r in results:
        if not r["gt_sku"]:
            continue
        if any(m["sku"] == r["gt_sku"] for m in r["top"][:3]):
            top3_hits += 1
        else:
            misses.append((r["title"], r["gt_sku"], [m["sku"] for m in r["top"][:3]]))
    print(f"QA GT test: {top3_hits}/{len(leads)} GT leads in top-3")
    for m in misses:
        print("   MISS:", m)
    if top3_hits < 18:
        sys.exit("QA FAIL: fewer than 18/21 GT leads in top-3 — fix the tables before shipping")

    # ---------------- workbook ----------------
    HEAD_FILL = PatternFill("solid", fgColor="D9D9D9")
    bold = Font(bold=True)
    wb = openpyxl.Workbook()

    ws = wb.active
    ws.title = "How to use"
    lines = [
        ("Resolve the 42 open cases — eval input titles vs catalogue (manual verification aid)", True),
        ("", False),
        ("The 42: input products from the brand-name eval run that the Power Query fuzzy merge left", False),
        ("below 0.70 AND that carry no embedded SKU code (the 58 located = 48 fuzzy + 10 code-rescued).", False),
        ("", False),
        ("Strictness (your decision, 2026-10-19): 'found' requires an EXACT variant. Where catalogue", False),
        ("variants differ on a dimension the title does not pin (length / class / colour / fit type),", False),
        ("the row is flagged AMBIG-VARIANT — do not count it as found unless you accept a variant call.", False),
        ("", False),
        ("Columns:", True),
        ("  GT lead SKU / name — the run export's recorded expected SKU (all 21 verified present in", False),
        ("    the 18-Oct catalogue export). A LEAD for your check, not a decision.", False),
        ("  Suggested status — GT-LEAD | EXACT candidate | AMBIG-VARIANT | NO-MATCH (script's read;", False),
        ("    your verdict is the verdict).", False),
        ("  Type/Size/Class/Length/Material — Y = the title's component matched the best candidate;", False),
        ("    N = present in the title but not matched; - = not applicable / absent.", False),
        ("  Variant gap — dimensions on which the candidate set differs but the title is silent.", False),
        ("  VERDICT (last column) — you fill: EXACT / AMBIG-VARIANT / NOT FOUND.", False),
        ("", False),
        ("Sheets: 'Resolve 42' (one row per input) | 'All candidates' (long list, rank order) |", False),
        ("'Not-found evidence' (token-level absence evidence for NO-MATCH rows).", False),
        ("", False),
        ("Scope note: only the 18-Oct catalogue export (2,335 products) is a source of truth here;", False),
        ("the 412xxx SKU family is absent from that export, so it cannot be 'found'.", False),
        ("No cell is auto-decided; the verdict column ships empty.", False),
    ]
    for i, (t, b) in enumerate(lines, 1):
        c = ws.cell(row=i, column=1, value=t)
        if b:
            c.font = bold
    ws.column_dimensions["A"].width = 115

    ws = wb.create_sheet("Resolve 42")
    hdr = ["#", "GT lead SKU", "GT lead catalogue name", "GT in export",
           "Suggested status", "Input product (from eval run)",
           "Best candidate SKU", "Best candidate catalogue name",
           "Type", "Size", "Class", "Length", "Material",
           "Variant gap", "Alt 2 SKU", "Alt 2 name", "Alt 3 SKU", "Alt 3 name",
           "VERDICT (EXACT / AMBIG-VARIANT / NOT FOUND)", "Note"]
    for j, h in enumerate(hdr, 1):
        c = ws.cell(row=1, column=j, value=h)
        c.font = bold
        c.fill = HEAD_FILL
        c.alignment = openpyxl.styles.Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(hdr))}{len(results)+1}"

    def gap_text(r):
        if r["suggestion"] == "GT-LEAD":
            return "; ".join(r["gaps"]) if r["gaps"] else ""
        return "; ".join(r["gaps"])

    for i, r in enumerate(results, 2):
        best = r["best"]
        alt2 = r["top"][1] if len(r["top"]) > 1 else None
        alt3 = r["top"][2] if len(r["top"]) > 2 else None
        note_bits = r["note_str"]
        row = [r["idx"], r["gt_sku"], r["gt_name"], r["gt_in"], r["suggestion"], r["title"],
               best["sku"] if best else "", best["name"] if best else "",
               r["flags"].get("type", ""), r["flags"].get("size", ""),
               r["flags"].get("class", ""), r["flags"].get("length", ""),
               r["flags"].get("material", ""),
               gap_text(r),
               alt2["sku"] if alt2 else "", alt2["name"] if alt2 else "",
               alt3["sku"] if alt3 else "", alt3["name"] if alt3 else "",
               "", note_bits]
        for j, v in enumerate(row, 1):
            c = ws.cell(row=i, column=j, value=v)
            if j in (2, 7, 15, 17):
                c.number_format = "@"
            if j in (3, 6, 8, 14, 16, 18, 20):
                c.alignment = openpyxl.styles.Alignment(wrap_text=True, vertical="top")
    for col, wd in {"A": 5, "B": 12, "C": 38, "D": 11, "E": 16, "F": 48, "G": 12, "H": 44,
                    "I": 6, "J": 6, "K": 6, "L": 7, "M": 8, "N": 40, "O": 12, "P": 40,
                    "Q": 12, "R": 40, "S": 22, "T": 40}.items():
        ws.column_dimensions[col].width = wd

    ws = wb.create_sheet("All candidates")
    hdr = ["#", "Input product (from eval run)", "Rank", "SKU", "Catalogue product name",
           "Type", "Size", "Class", "Length", "Material", "Score"]
    for j, h in enumerate(hdr, 1):
        c = ws.cell(row=1, column=j, value=h)
        c.font = bold
        c.fill = HEAD_FILL
    for i, m in enumerate(cand_long, 2):
        row = [m["idx"], m["title"], m["rank"], m["sku"], m["name"],
               m["type"], m["size"], m["class"], m["length"], m["material"], m["score"]]
        for j, v in enumerate(row, 1):
            c = ws.cell(row=i, column=j, value=v)
            if j == 4:
                c.number_format = "@"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:K{len(cand_long)+1}"
    for col, wd in {"A": 5, "B": 46, "C": 6, "D": 12, "E": 52, "F": 6, "G": 6,
                    "H": 6, "I": 7, "J": 8, "K": 8}.items():
        ws.column_dimensions[col].width = wd

    ws = wb.create_sheet("Not-found evidence")
    hdr = ["#", "Input product", "Concept", "Tokens absent from catalogue",
           "Nearest family (up to 3)"]
    for j, h in enumerate(hdr, 1):
        c = ws.cell(row=1, column=j, value=h)
        c.font = bold
        c.fill = HEAD_FILL
    for i, r in enumerate(no_match, 2):
        ip = parse(r["title"])
        near = []
        for sku, name, cp in catp:
            if ip["sizes"] and not size_ok(ip["sizes"], cp["sizes"]):
                continue
            if r["concept"] not in (None, "UNKNOWN"):
                if not candidate_type_match(r["concept"], norm(name), cp):
                    continue
            near.append((sku, name))
            if len(near) == 3:
                break
        row = [r["idx"], r["title"], r["concept"] or "-",
               ", ".join(r["zero_tokens"]) or "-",
               "; ".join(f"{s} {n}" for s, n in near)]
        for j, v in enumerate(row, 1):
            c = ws.cell(row=i, column=j, value=v)
            if j in (2, 5):
                c.alignment = openpyxl.styles.Alignment(wrap_text=True, vertical="top")
    for col, wd in {"A": 5, "B": 46, "C": 14, "D": 30, "E": 90}.items():
        ws.column_dimensions[col].width = wd

    xlsx = outdir / "open_42_resolution.xlsx"
    wb.save(xlsx)
    print("saved:", xlsx)

    # ---------------- CSV sidecars ----------------
    wide = outdir / "open_42_resolution.csv"
    with open(wide, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["#", "gt_lead_sku", "gt_lead_name", "gt_in_export", "suggested_status",
                    "input_product", "best_sku", "best_name", "type", "size", "class",
                    "length", "material", "variant_gap", "alt2_sku", "alt2_name",
                    "alt3_sku", "alt3_name", "verdict", "note"])
        for r in results:
            best = r["best"]
            alt2 = r["top"][1] if len(r["top"]) > 1 else None
            alt3 = r["top"][2] if len(r["top"]) > 2 else None
            w.writerow([r["idx"], r["gt_sku"], r["gt_name"], r["gt_in"], r["suggestion"], r["title"],
                        best["sku"] if best else "", best["name"] if best else "",
                        r["flags"].get("type", ""), r["flags"].get("size", ""),
                        r["flags"].get("class", ""), r["flags"].get("length", ""),
                        r["flags"].get("material", ""), "; ".join(r["gaps"]),
                        alt2["sku"] if alt2 else "", alt2["name"] if alt2 else "",
                        alt3["sku"] if alt3 else "", alt3["name"] if alt3 else "", "", r["note_str"]])
    print("saved:", wide)

    long_csv = outdir / "open_42_candidates_long.csv"
    with open(long_csv, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["#", "input_product", "rank", "sku", "name", "type", "size", "class",
                    "length", "material", "score"])
        for m in cand_long:
            w.writerow([m["idx"], m["title"], m["rank"], m["sku"], m["name"], m["type"],
                        m["size"], m["class"], m["length"], m["material"], m["score"]])
    print("saved:", long_csv)

    manifest = outdir / "open_cases_42.csv"
    with open(manifest, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["#", "input_product", "concept", "gt_status", "gt_sku"])
        for r in results:
            w.writerow([r["idx"], r["title"], r["concept"] or "", r["gt_status"], r["gt_sku"]])
    print("saved:", manifest)

    # ---------------- summary + copy ----------------
    from collections import Counter
    cnt = Counter(r["suggestion"] for r in results)
    print("suggestions:", dict(cnt))
    print("no-match rows:", [(r["idx"], r["title"]) for r in no_match])

    if args.copy_to:
        dest = Path(args.copy_to) / xlsx.name
        shutil.copy2(xlsx, dest)
        print("copied to:", dest)

    # read-back QA
    wb2 = openpyxl.load_workbook(xlsx, read_only=True)
    print("QA sheets:", wb2.sheetnames)
    ws2 = wb2["Resolve 42"]
    print("QA Resolve 42 rows (incl header):", ws2.max_row)
    verdicts = [r[18] for r in ws2.iter_rows(min_row=2, values_only=True)]
    print("QA verdict column empty:", all(v in (None, "") for v in verdicts))
    wb2.close()


if __name__ == "__main__":
    main()
