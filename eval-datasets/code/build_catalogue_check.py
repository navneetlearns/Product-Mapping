#!/usr/bin/env python3
"""Build a VLOOKUP-ready workbook: eval-run INPUT products vs the product catalogue.

Purpose: take the eval run's input products (the invoice product titles the
brand-name eval run sent in — one per case) and lay them out as a clean column
so they can be compared, product per product, against the catalogue's
Product Name column — manually / by VLOOKUP.
This script only CLEANS and lays out the data; it computes no coverage counts.

Default sources (the two files in the F: drop folder):
  - run       : products-.../Eval Related Product Mapping/turnover-product-mapping-c4cce19d.csv
  - catalogue : products-.../Eval Related Product Mapping/products-18-October-2026-10-52-16.csv

Output workbook sheets: How to use | Input products | Catalogue | Catalogue A-Z
All SKU/code columns are written as TEXT so VLOOKUPs stay type-consistent.

Usage:
  python3 build_catalogue_check.py [--run <run.csv>] [--catalogue <cat.csv>] [--copy-to <dir>]
"""
import argparse
import csv
import json
import re
import unicodedata
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
OUTDIR = HERE.parent / "outputs" / "catalogue_check"

DROP = Path("/mnt/f/Product/zotok/Product Mapping/Eval Related Product Mapping")
DEF_RUN = DROP / "turnover-product-mapping-c4cce19d.csv"
DEF_CAT = DROP / "products-18-October-2026-10-52-16.csv"


def _to_bytes(s: str) -> bytes:
    """Re-encode a mojibake string back to its original bytes.

    cp1252 for everything it defines (so '€' -> 0x80), raw ord() for the
    C1 range it leaves undefined (0x80-0x9F like '\\x9d'), which is exactly
    where utf-8 continuation bytes land in this class of mojibake."""
    out = bytearray()
    for ch in s:
        try:
            out += ch.encode("cp1252")
        except UnicodeEncodeError:
            if ord(ch) <= 0xFF:
                out.append(ord(ch))
            else:
                raise
    return bytes(out)


def clean(s) -> str:
    """Repair mojibake, strip zero-widths/C1 controls, collapse whitespace."""
    if s is None:
        return ""
    s = str(s)
    if any(ch in s for ch in ("Â", "â", "€", "Ã")):
        try:
            s = _to_bytes(s).decode("utf-8")
        except (UnicodeEncodeError, UnicodeDecodeError):
            pass
    s = "".join(ch for ch in s if unicodedata.category(ch) != "Cf" and not (0x80 <= ord(ch) <= 0x9F))
    s = s.replace("\xa0", " ").replace("\u200b", "")
    return " ".join(s.split())


def sku_sort_key(s):
    return (len(s), s)


def load_run_inputs(path):
    """Unique input products in first-appearance order (hang+rerun pairs collapse)."""
    rows = list(csv.DictReader(open(path, newline="", encoding="utf-8")))
    order, info, collapsed = [], {}, []
    for r in rows:
        li = json.loads(r["inputs"])["line_items"][0]
        t = li["product_name"]
        if t in info:
            collapsed.append((t, info[t]["status"], r["status"]))
            continue
        info[t] = {"status": r["status"], "rep": r["repetition"]}
        order.append(t)
    return rows, order, info, collapsed


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--run", default=str(DEF_RUN))
    ap.add_argument("--catalogue", default=str(DEF_CAT))
    ap.add_argument("--copy-to")
    args = ap.parse_args()

    OUTDIR.mkdir(parents=True, exist_ok=True)

    # ---------------- eval run inputs ----------------
    run_rows, order, info, collapsed = load_run_inputs(args.run)
    inputs = [clean(t) for t in order]
    print(f"run: {len(run_rows)} export rows -> {len(inputs)} unique input products; "
          f"collapsed duplicates: {len(collapsed)}")
    for c in collapsed:
        print("   collapsed (hang+rerun):", c)
    moji_in = [t for t in order if any(ch in t for ch in ("Â", "â", "Ã", "€", "\x9d", "\u200b", "\xa0"))]
    print("   inputs with mojibake/zero-width markers:", len(moji_in))
    for t in moji_in[:5]:
        print("      cleaned:", repr(t), "->", repr(clean(t)))

    # ---------------- catalogue ----------------
    raw = list(csv.DictReader(open(args.catalogue, newline="", encoding="utf-8")))
    cat_rows, seen = [], {}
    for r in raw:
        name = clean(r.get("Product Name"))
        sku = clean(r.get("Sku Code"))
        if sku in seen:
            print(f"WARN duplicate sku {sku} -> keeping first")
            continue
        seen[sku] = name
        cat_rows.append((sku, name))
    cat_rows = sorted(cat_rows, key=lambda t: sku_sort_key(t[0]))
    moji_cat = sum(1 for s, n in cat_rows if any(ch in n for ch in ("Â", "â", "Ã", "€")))
    print(f"catalogue: {len(raw)} raw rows -> {len(cat_rows)} clean rows; mojibake remaining: {moji_cat}")

    # ---------------- workbook ----------------
    wb = openpyxl.Workbook()
    HEAD_FILL = PatternFill("solid", fgColor="D9D9D9")
    bold = Font(bold=True)

    ws = wb.active
    ws.title = "How to use"
    run_name = Path(args.run).stem
    lines = [
        ("Eval run inputs vs Catalogue — for manual one-to-one comparison", True),
        ("", False),
        (f"Input products: the 100 products sent into the brand-name eval run '{run_name}'", False),
        ("(LangSmith experiment; the invoice product titles as they entered the mapping pipeline).", False),
        ("  The export has 102 rows -> 100 unique products: 2 titles appear twice because they hung", False),
        ("  and were rerun (same product, kept once; the 100-row list is the deduped set).", False),
        ("", False),
        ("Compare against: the 'Catalogue' sheet's Product Name column.", False),
        ("  Catalogue = products-18-October-2026 export (2,335 products), cleaned: whitespace/", False),
        ("  zero-width/mojibake fixed, duplicate Description column dropped, no duplicate SKUs,", False),
        ("  sorted by Sku Code. Both sheets' text was cleaned the same way.", False),
        ("", False),
        ("How to compare: for each input product, find its product in the Catalogue sheet and", False),
        ("record it next to the input (add columns B/C as needed). None of these 100 titles match", False),
        ("a catalogue name exactly — invoice titles differ from catalogue names (brand tokens,", False),
        ("word order, quotes) — so scan by the key words (e.g. 'UPVC TEE 1 1/2') or use Excel's", False),
        ("fuzzy lookup (Power Query).", False),
        ("Formulas (A2 = an input product):", False),
        ("  nearest-in-A-Z guide:  =VLOOKUP(A2, 'Catalogue A-Z'!$A$2:$B$2336, 2, TRUE)", False),
        ("     TRUE = lower-bound: returns the SKU of the last catalogue name alphabetically", False),
        ("     <= A2 - jump to that spot in the A-Z list and scan around it. It is NOT a", False),
        ("     similarity match, and it NEEDS the ascending sort (why 'Catalogue A-Z' exists).", False),
        ("  exact match only:      =VLOOKUP(A2, Catalogue!$B$2:$B$2336, 1, FALSE)", False),
        ("     The 3rd argument must be a NUMBER (2 = 2nd column of the table), never C:C.", False),
        ("All SKU/code columns are stored as TEXT — keep them text in your lookups.", False),
    ]
    for i, (t, b) in enumerate(lines, 1):
        c = ws.cell(row=i, column=1, value=t)
        if b:
            c.font = bold
    ws.column_dimensions["A"].width = 110

    ws = wb.create_sheet("Input products")
    c = ws.cell(row=1, column=1, value="Input Product (from eval run)")
    c.font = bold
    c.fill = HEAD_FILL
    for i, t in enumerate(inputs, 2):
        ws.cell(row=i, column=1, value=t)
    ws.column_dimensions["A"].width = 70
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:A{len(inputs)+1}"

    ws = wb.create_sheet("Catalogue")
    for j, h in enumerate(["Sku Code", "Product Name"], 1):
        c = ws.cell(row=1, column=j, value=h)
        c.font = bold
        c.fill = HEAD_FILL
    for i, (sku, name) in enumerate(cat_rows, 2):
        ws.cell(row=i, column=1, value=sku).number_format = "@"
        ws.cell(row=i, column=2, value=name)
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 70
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:B{len(cat_rows)+1}"

    ws = wb.create_sheet("Catalogue A-Z")
    for j, h in enumerate(["Product Name", "Sku Code"], 1):
        c = ws.cell(row=1, column=j, value=h)
        c.font = bold
        c.fill = HEAD_FILL
    az = sorted(cat_rows, key=lambda t: (t[1].casefold(), t[1]))
    for i, (sku, name) in enumerate(az, 2):
        ws.cell(row=i, column=1, value=name)
        ws.cell(row=i, column=2, value=sku).number_format = "@"
    ws.column_dimensions["A"].width = 70
    ws.column_dimensions["B"].width = 14
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:B{len(az)+1}"

    xlsx = OUTDIR / "eval_inputs_vs_catalogue_vlookup.xlsx"
    wb.save(xlsx)
    print("saved:", xlsx)

    csv_path = OUTDIR / "eval_run_input_products.csv"
    with open(csv_path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["Input Product (from eval run)"])
        for t in inputs:
            w.writerow([t])
    print("saved:", csv_path)

    cat_csv = OUTDIR / "catalogue_cleaned.csv"
    with open(cat_csv, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["Sku Code", "Product Name"])
        w.writerows(cat_rows)
    print("saved:", cat_csv)

    if args.copy_to:
        import shutil
        dest = Path(args.copy_to) / xlsx.name
        shutil.copy2(xlsx, dest)
        print("copied to:", dest)

    # ---------------- QA ----------------
    wb2 = openpyxl.load_workbook(xlsx, read_only=True)
    print("QA Input products rows:", wb2["Input products"].max_row - 1)
    print("QA Catalogue rows:", wb2["Catalogue"].max_row - 1)
    print("QA Catalogue A-Z rows:", wb2["Catalogue A-Z"].max_row - 1)
    wb2.close()


if __name__ == "__main__":
    main()
